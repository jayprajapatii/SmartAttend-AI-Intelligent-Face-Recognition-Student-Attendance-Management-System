"""
reports/excel_report_generator.py

Generates styled .xlsx reports via openpyxl: bold colored header row,
auto-sized columns, frozen header row, and an optional summary block.
This is the reusable version of the export logic currently inline in
ui/pages/reports_page.py's _export_xlsx() - swap that page to call
ExcelReportGenerator.generate() instead.

Usage:
    from reports.excel_report_generator import ExcelReportGenerator

    ExcelReportGenerator.generate(
        title="Daily Attendance Report - 2026-08-09",
        columns=[("roll_number", "Roll No."), ("full_name", "Name"), ("attendance_status", "Status")],
        rows=[{"roll_number": "101", "full_name": "Jay Prajapati", "attendance_status": "Present"}],
        filepath="data/reports/daily_attendance.xlsx",
    )

    # With a summary section (e.g. counts by status):
    ExcelReportGenerator.generate(
        title="...", columns=[...], rows=[...], filepath="...",
        summary={"Present": 28, "Absent": 4, "Late": 2},
    )
"""

import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("SmartAttendAI.reports.excel_report_generator")

HEADER_FILL_COLOR = "3B82F6"
HEADER_FONT_COLOR = "FFFFFF"
TITLE_FONT_SIZE = 14
MIN_COLUMN_WIDTH = 12
MAX_COLUMN_WIDTH = 50


class ExcelReportGenerator:
    """Static-method generator - no instance state needed for Excel report creation."""

    @staticmethod
    def generate(
        title: str,
        columns: list,
        rows: list,
        filepath: str,
        summary: dict = None,
        sheet_name: str = "Report",
    ) -> str:
        """
        columns: list of (key, header) tuples.
        rows: list of dicts, each expected to have the column keys.
        summary: optional dict of label -> value, written to a small
                 block above the data table (e.g. {"Present": 28, "Absent": 4}).
        Returns the filepath written to.
        """
        filepath = Path(filepath)
        filepath.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        current_row = 1
        current_row = ExcelReportGenerator._write_title(ws, title, current_row, len(columns))

        if summary:
            current_row = ExcelReportGenerator._write_summary(ws, summary, current_row)

        header_row = current_row
        ExcelReportGenerator._write_header(ws, columns, header_row)

        data_start_row = header_row + 1
        ExcelReportGenerator._write_rows(ws, columns, rows, data_start_row)

        ExcelReportGenerator._auto_size_columns(ws, columns, header_row)
        ws.freeze_panes = ws.cell(row=data_start_row, column=1)

        wb.save(filepath)
        logger.info("Excel report written: %s (%d row(s)).", filepath, len(rows))
        return str(filepath)

    # ------------------------------------------------------------
    # Section writers
    # ------------------------------------------------------------
    @staticmethod
    def _write_title(ws, title: str, row: int, column_span: int) -> int:
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = Font(size=TITLE_FONT_SIZE, bold=True)
        if column_span > 1:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=column_span)
        return row + 2  # one blank row after the title

    @staticmethod
    def _write_summary(ws, summary: dict, start_row: int) -> int:
        row = start_row
        for label, value in summary.items():
            label_cell = ws.cell(row=row, column=1, value=f"{label}:")
            label_cell.font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)
            row += 1
        return row + 1  # one blank row after the summary block

    @staticmethod
    def _write_header(ws, columns: list, row: int):
        header_fill = PatternFill(start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type="solid")
        header_font = Font(bold=True, color=HEADER_FONT_COLOR)
        thin_border = Border(bottom=Side(style="thin", color="CCCCCC"))

        for col_idx, (_key, header) in enumerate(columns, start=1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = thin_border

    @staticmethod
    def _write_rows(ws, columns: list, rows: list, start_row: int):
        keys = [c[0] for c in columns]
        stripe_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

        for row_offset, row_data in enumerate(rows):
            excel_row = start_row + row_offset
            is_striped = row_offset % 2 == 1

            for col_idx, key in enumerate(keys, start=1):
                value = row_data.get(key)
                cell = ws.cell(row=excel_row, column=col_idx, value=value if value is not None else "-")
                if is_striped:
                    cell.fill = stripe_fill

    @staticmethod
    def _auto_size_columns(ws, columns: list, header_row: int):
        for col_idx, (_key, header) in enumerate(columns, start=1):
            column_letter = get_column_letter(col_idx)
            max_length = len(str(header))

            for row in ws.iter_rows(min_row=header_row + 1, min_col=col_idx, max_col=col_idx):
                for cell in row:
                    if cell.value is not None:
                        max_length = max(max_length, len(str(cell.value)))

            ws.column_dimensions[column_letter].width = min(
                max(max_length + 2, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH
            )